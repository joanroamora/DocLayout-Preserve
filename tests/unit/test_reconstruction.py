"""Unit tests for AdaptiveTextFitter, PDFDocumentBuilder, and DocxDocumentBuilder."""

from __future__ import annotations

import base64
import io

import docx
import pymupdf
import pytest
from PIL import Image

from doclayout.core.models import (
    BlockType,
    BoundingBox,
    DocumentLayout,
    LayoutBlock,
    PageLayout,
    TableCell,
    TransformationResult,
    TransformationStatus,
)
from doclayout.ingestion import extract_document_layout
from doclayout.llm import TaskType, TransformationRequest
from doclayout.llm.schemas import DocumentTransformationOutput
from doclayout.reconstruction import (
    AdaptiveTextFitter,
    reconstruct_document,
)
from tests.mocks import MockGeminiClient


def _make_tiny_png_b64() -> str:
    """Create a 16x16 solid blue PNG image encoded in base64."""
    img = Image.new("RGB", (16, 16), color=(30, 90, 200))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("ascii")


def test_adaptive_text_fitter_natural_fit() -> None:
    """Verify short text fits naturally with zero scaling or negative tracking."""
    fitter = AdaptiveTextFitter(min_font_scale_ratio=0.75)
    bbox = BoundingBox(x0=50.0, y0=50.0, x1=350.0, y1=90.0, page_index=0)

    metrics = fitter.fit_text_to_box("Short headline", bbox=bbox, base_font_size=12.0)
    assert metrics.fitted_font_size == pytest.approx(12.0)
    assert metrics.applied_tracking == pytest.approx(0.0)
    assert metrics.applied_line_height == pytest.approx(1.2)
    assert metrics.overflow_detected is False


def test_adaptive_text_fitter_micro_tracking_only() -> None:
    """Verify borderline line-wrap is resolved by Adjustment 1 (micro-tracking) at 100% font size."""
    fitter = AdaptiveTextFitter(min_font_scale_ratio=0.75)
    # Measure exact width of sample string at 12pt Helvetica and make box ~2% narrower
    sample = " Architecture Preserving Specification"
    exact_w = float(pymupdf.get_text_length(sample.strip(), fontname="helv", fontsize=12.0))
    tight_width = exact_w * 0.98
    # Height allows 1 line (14.4pt) -> height=16pt, so wrapping to 2 lines (28.8pt) wouldn't fit
    bbox = BoundingBox(x0=10.0, y0=10.0, x1=10.0 + tight_width, y1=26.0, page_index=0)

    metrics = fitter.fit_text_to_box(sample, bbox=bbox, base_font_size=12.0)
    assert metrics.fitted_font_size == pytest.approx(12.0)
    assert metrics.applied_tracking < 0.0
    assert metrics.overflow_detected is False


def test_adaptive_text_fitter_proportional_font_scaling_and_overflow() -> None:
    """Verify proportional scaling down to 75% and overflow warning when exceeded."""
    fitter = AdaptiveTextFitter(min_font_scale_ratio=0.75)
    bbox = BoundingBox(x0=50.0, y0=50.0, x1=250.0, y1=85.0, page_index=0)

    moderate_text = (
        "Este párrafo traducido ocupa ligeramente más espacio que el bloque en inglés "
        "original pero encaja reduciendo proporcionalmente la fuente."
    )
    scaled = fitter.fit_text_to_box(moderate_text, bbox=bbox, base_font_size=12.0)
    assert 9.0 <= scaled.fitted_font_size < 12.0
    assert scaled.overflow_detected is False

    # Massive text that cannot fit even at 75% (9.0pt) inside 200x35pt box
    huge_text = " ".join(["Texto demasiado extenso para el contenedor físico."] * 12)
    block = LayoutBlock(block_id="blk_over", bbox=bbox, text="Short", font_size=12.0)
    base_res = TransformationResult(
        block_id="blk_over",
        original_text="Short",
        transformed_text=huge_text,
    )
    enriched = fitter.apply_to_result(block, base_res)
    assert enriched.applied_font_size == pytest.approx(9.0)
    assert enriched.overflow_detected is True
    assert enriched.status == TransformationStatus.OVERFLOW_WARNING
    assert any("Desbordamiento tipográfico" in w for w in enriched.warnings)


def test_pdf_overlay_reconstruction_preserves_vectors_and_replaces_text() -> None:
    """Verify overlay PDF reconstruction strips old text, keeps vector lines, and inserts new text."""
    src_doc = pymupdf.open()
    p = src_doc.new_page(width=595, height=842)
    p.insert_text((50, 90), "Original Architecture Title", fontsize=16)
    p.insert_text((50, 140), "Body paragraph to translate cleanly.", fontsize=11)

    # Decorative vector line + 2x2 table grid
    p.draw_line((50, 110), (500, 110))
    for y in (220, 250, 280):
        p.draw_line((50, y), (350, y))
    for x in (50, 200, 350):
        p.draw_line((x, 220), (x, 280))
    p.insert_text((60, 240), "ColA", fontsize=10)
    p.insert_text((210, 240), "ColB", fontsize=10)
    p.insert_text((60, 270), "Val1", fontsize=10)
    p.insert_text((210, 270), "Val2", fontsize=10)

    original_bytes = bytes(src_doc.tobytes())
    original_drawings_count = len(p.get_drawings())
    src_doc.close()

    layout = extract_document_layout(original_bytes, filename="overlay.pdf")
    tx_out = MockGeminiClient().transform_document(
        layout,
        TransformationRequest(task_type=TaskType.TRANSLATE, target_language="Español"),
    )

    artifact = reconstruct_document(
        layout=layout,
        transformation=tx_out,
        output_format="pdf",
        original_pdf_bytes=original_bytes,
    )

    assert artifact.output_format == "pdf"
    assert len(artifact.output_bytes) > 0

    out_doc = pymupdf.open(stream=artifact.output_bytes, filetype="pdf")
    try:
        out_page = out_doc[0]
        normalized_text = " ".join(out_page.get_text().split())
        assert "[ES] Original Architecture Title" in normalized_text
        assert "[ES] Body paragraph to translate cleanly." in normalized_text
        assert "[ES] ColA" in normalized_text
        assert len(out_page.get_drawings()) >= original_drawings_count
    finally:
        out_doc.close()


def test_standalone_pdf_and_docx_reconstruction_with_images_and_tables() -> None:
    """Verify standalone PDF and DOCX builders render headings, paragraphs, images, and tables."""
    png_b64 = _make_tiny_png_b64()
    h_blk = LayoutBlock(
        block_id="blk_h",
        block_type=BlockType.HEADING,
        bbox=BoundingBox(x0=50.0, y0=50.0, x1=500.0, y1=85.0, page_index=0),
        text="Section Title",
        font_size=16.0,
        color="#1A3B5C",
    )
    t_blk = LayoutBlock(
        block_id="blk_t",
        block_type=BlockType.TABLE,
        bbox=BoundingBox(x0=50.0, y0=100.0, x1=450.0, y1=160.0, page_index=0),
        text="Header1 | Header2",
        cells=[
            TableCell(row_index=0, col_index=0, text="Header1"),
            TableCell(row_index=0, col_index=1, text="Header2"),
        ],
        metadata={"row_count": 1, "col_count": 2, "relative_col_widths": [0.5, 0.5]},
    )
    img_blk = LayoutBlock(
        block_id="blk_i",
        block_type=BlockType.IMAGE,
        bbox=BoundingBox(x0=50.0, y0=180.0, x1=130.0, y1=260.0, page_index=0),
        image_bytes_b64=png_b64,
    )
    layout = DocumentLayout(
        document_id="doc_multi",
        filename="report.docx",
        source_format="docx",
        pages=[PageLayout(page_index=0, blocks=[h_blk, t_blk, img_blk])],
    )

    tx_out: DocumentTransformationOutput = MockGeminiClient().transform_document(
        layout,
        TransformationRequest(task_type=TaskType.TRANSLATE, target_language="Francés"),
    )

    # 1. Standalone PDF export
    pdf_artifact = reconstruct_document(layout=layout, transformation=tx_out, output_format="pdf")
    pdf_doc = pymupdf.open(stream=pdf_artifact.output_bytes, filetype="pdf")
    try:
        assert "[FR] Section Title" in pdf_doc[0].get_text()
        assert len(pdf_doc[0].get_images()) == 1
    finally:
        pdf_doc.close()

    # 2. Native DOCX export
    docx_artifact = reconstruct_document(layout=layout, transformation=tx_out, output_format="docx")
    reopened_docx = docx.Document(io.BytesIO(docx_artifact.output_bytes))
    all_para_text = "\n".join(p.text for p in reopened_docx.paragraphs)
    assert "[FR] Section Title" in all_para_text
    assert len(reopened_docx.tables) == 1
    assert reopened_docx.tables[0].cell(0, 0).text == "[FR] Header1"
    assert len(reopened_docx.inline_shapes) == 1

    with pytest.raises(ValueError):
        reconstruct_document(layout=layout, transformation=tx_out, output_format="html")
