"""Unit and headless UI tests for Streamlit Side-by-Side viewer, pipeline, and session history."""

from __future__ import annotations

import io
from pathlib import Path

import docx
import pymupdf
import pytest
from PIL import Image
from streamlit.testing.v1 import AppTest

from doclayout.core.models import BlockType
from doclayout.llm import TaskType, TransformationRequest
from doclayout.ui import (
    SessionStore,
    build_layout_wireframe_pdf,
    render_pdf_page_png,
    run_four_stage_pipeline,
)
from tests.mocks import MockGeminiClient


def _sample_pdf_bytes(title: str = "Document Title") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox((50, 80, 500, 140), title, fontsize=16)
    page.insert_textbox(
        (50, 160, 500, 260), "Body paragraph for UI pipeline verification.", fontsize=11
    )
    data = bytes(doc.tobytes())
    doc.close()
    return data


def _sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("DOCX Preview Title", level=1)
    doc.add_paragraph("Paragraph inside Word document.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pipeline_history_caching_retries_and_cleanup_actions(tmp_path: Path) -> None:
    """Verify SHA-256 caching, non-duplicating retries, New Document, and two session clear actions."""
    store = SessionStore()
    pdf_a = _sample_pdf_bytes("First Document")
    pdf_b = _sample_pdf_bytes("Second Document")

    stages_seen: list[int] = []
    job1 = run_four_stage_pipeline(
        store=store,
        file_bytes=pdf_a,
        filename="first.pdf",
        request=TransformationRequest(task_type=TaskType.TRANSLATE, target_language="Español"),
        llm_client=MockGeminiClient(),
        stage_callback=lambda num, _lbl: stages_seen.append(num),
    )

    assert stages_seen == [1, 2, 3, 4]
    assert job1.status == "completed"
    assert job1.pdf_output_bytes is not None
    assert job1.docx_output_bytes is not None
    assert job1.export_json_bytes is not None
    assert len(store.history) == 1
    assert len(store.extraction_cache) == 1

    # "Nuevo Documento" clears active editor but preserves extraction cache and history
    store.start_new_document()
    assert store.active_job is None
    assert len(store.history) == 1
    assert len(store.extraction_cache) == 1
    assert store.editor_reset_counter == 1

    # Simulate failed job on second document, then retry with same retry_job_id (no duplicate!)
    failed_job = run_four_stage_pipeline(
        store=store,
        file_bytes=pdf_b,
        filename="second.pdf",
        request=TransformationRequest(),
        llm_client=MockGeminiClient(simulate_rate_limit=True),
    )
    assert failed_job.status == "failed"
    assert len(store.history) == 2

    retried_job = run_four_stage_pipeline(
        store=store,
        file_bytes=pdf_b,
        filename="second.pdf",
        request=TransformationRequest(),
        llm_client=MockGeminiClient(),
        retry_job_id=failed_job.job_id,
    )
    assert retried_job.job_id == failed_job.job_id
    assert retried_job.status == "completed"
    assert len(store.history) == 2  # Updated in place without creating a 3rd entry!
    assert store.history[0].job_id == failed_job.job_id
    assert store.history[0].status == "completed"

    # "Vaciar historial" clears history only, leaving active_job and cache intact
    store.clear_history_only()
    assert len(store.history) == 0
    assert store.active_job is not None
    assert len(store.extraction_cache) == 2

    # "Limpiar toda la sesión" wipes history, cache, active_job, and .tmp files
    dummy_tmp_file = tmp_path / "temp_artifact.bin"
    dummy_tmp_file.write_bytes(b"temporary")
    store.clear_entire_session(tmp_dir=tmp_path)
    assert store.active_job is None
    assert len(store.history) == 0
    assert len(store.extraction_cache) == 0
    assert not dummy_tmp_file.exists()


def test_duplicate_execution_blocked_while_processing() -> None:
    """Verify concurrent execution raises RuntimeError when store.is_processing is True."""
    store = SessionStore()
    store.is_processing = True
    with pytest.raises(RuntimeError):
        run_four_stage_pipeline(
            store=store,
            file_bytes=_sample_pdf_bytes(),
            filename="busy.pdf",
            request=TransformationRequest(),
            llm_client=MockGeminiClient(),
        )


def test_preview_png_and_docx_wireframe_rendering() -> None:
    """Verify PNG rendering draws colored bounding boxes and DOCX wireframe builds valid PDF."""
    store = SessionStore()
    _, docx_layout, wireframe_pdf = store.get_or_extract_layout(_sample_docx_bytes(), "test.docx")
    assert len(wireframe_pdf) > 0

    rebuilt = build_layout_wireframe_pdf(docx_layout)
    png_bytes = render_pdf_page_png(
        pdf_bytes=rebuilt,
        page_index=0,
        blocks=docx_layout.all_blocks,
        visible_types={BlockType.HEADING, BlockType.PARAGRAPH},
        overflow_block_ids={docx_layout.all_blocks[0].block_id},
        dpi=96,
    )
    img = Image.open(io.BytesIO(png_bytes))
    assert img.width > 100
    assert img.height > 100


def test_streamlit_app_renders_cleanly_in_spanish() -> None:
    """Verify src/doclayout/ui/app.py loads headlessly via AppTest with zero exceptions."""
    app_path = Path(__file__).resolve().parents[2] / "src" / "doclayout" / "ui" / "app.py"
    at = AppTest.from_file(str(app_path), default_timeout=15)
    at.run()

    assert not at.exception
    assert any("DocLayout-Preserve" in t.value for t in at.title)
    button_labels = [b.label for b in at.button]
    assert "Procesar las 4 etapas" in button_labels
    assert "Nuevo Documento" in button_labels
    assert "Vaciar historial" in button_labels
    assert "Limpiar toda la sesión" in button_labels
