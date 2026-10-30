"""Unit and integration tests for Feature 8: Async REST API, Worker, Encrypted Store, and K8s."""

from __future__ import annotations

import json
import socket
import urllib.request
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

import pymupdf
import yaml
from docx import Document
from fastapi.testclient import TestClient

from doclayout.api import (
    AsyncDocumentWorker,
    EncryptedArtifactStore,
    InMemoryJobQueue,
    JobStage,
    JobState,
    SlidingWindowRateLimiter,
    WorkerProbeServer,
    ZeroQuotaSimulationClient,
    create_app,
)
from doclayout.core.config import Settings
from doclayout.core.models import DocumentLayout
from doclayout.llm.schemas import DocumentTransformationOutput, TransformationRequest
from tests.mocks.gemini_mock import MockGeminiClient

REPO_ROOT = Path(__file__).resolve().parents[2]


def _make_sample_pdf(secret_marker: str = "CONFIDENTIAL_LAYOUT_PAYLOAD_2026") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595.0, height=842.0)
    page.insert_text((72.0, 90.0), "1. Executive Summary", fontsize=16.0)
    page.insert_text((72.0, 130.0), f"Sensitive marker: {secret_marker}", fontsize=11.0)
    data = bytes(doc.tobytes())
    doc.close()
    return data


def _make_sample_docx() -> bytes:
    doc = Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("Asynchronous workers preserve layout fidelity across replicas.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_test_stack(
    tmp_path: Path,
    max_concurrent: int = 5,
    max_file_mb: int = 20,
    ttl_seconds: int = 3600,
) -> tuple[TestClient, EncryptedArtifactStore, InMemoryJobQueue, AsyncDocumentWorker]:
    cfg = Settings(
        ARTIFACT_STORE_DIR=str(tmp_path / "artifacts"),
        ARTIFACT_TTL_SECONDS=ttl_seconds,
        MAX_CONCURRENT_JOBS_PER_SESSION=max_concurrent,
        MAX_FILE_SIZE_MB=max_file_mb,
    )
    store = EncryptedArtifactStore(
        root_dir=tmp_path / "artifacts",
        ttl_seconds=ttl_seconds,
        encryption_secret="unit-test-encryption-secret",
    )
    queue = InMemoryJobQueue()
    worker = AsyncDocumentWorker(
        store=store,
        queue=queue,
        llm_client=MockGeminiClient(),
    )
    client = TestClient(create_app(settings=cfg, store=store, queue=queue))
    return client, store, queue, worker


def test_health_and_readiness_zero_quota(tmp_path: Path) -> None:
    client, _, _, _ = _build_test_stack(tmp_path)

    resp_h = client.get("/healthz")
    assert resp_h.status_code == 200
    assert resp_h.json()["status"] == "ok"

    resp_r = client.get("/readyz")
    assert resp_r.status_code == 200
    body = resp_r.json()
    assert body["status"] == "ready"
    assert body["storage_ready"] is True
    assert body["queue_ready"] is True
    assert body["queue_depth"] == 0


def test_magic_bytes_and_validation_guards(tmp_path: Path) -> None:
    client, _, _, _ = _build_test_stack(tmp_path, max_file_mb=1)

    # Missing X-Session-ID header
    r_no_sess = client.post(
        "/api/v1/jobs",
        files={"file": ("doc.pdf", _make_sample_pdf(), "application/pdf")},
    )
    assert r_no_sess.status_code == 400

    # Invalid extension
    r_bad_ext = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_alpha"},
        files={"file": ("malware.exe", b"MZ\x90\x00", "application/octet-stream")},
    )
    assert r_bad_ext.status_code == 422

    # Fake PDF without %PDF- magic header
    r_fake_pdf = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_alpha"},
        files={"file": ("fake.pdf", b"NOT_A_REAL_PDF_HEADER", "application/pdf")},
    )
    assert r_fake_pdf.status_code == 422

    # Fake DOCX without PK\x03\x04 ZIP header
    r_fake_docx = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_alpha"},
        files={"file": ("fake.docx", b"NOT_A_ZIP_CONTAINER", "application/octet-stream")},
    )
    assert r_fake_docx.status_code == 422

    # Oversized payload (> 1 MB configured limit)
    oversized = b"%PDF-1.7\n" + (b"0" * (1024 * 1024 + 32))
    r_large = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_alpha"},
        files={"file": ("huge.pdf", oversized, "application/pdf")},
    )
    assert r_large.status_code == 413


def test_end_to_end_async_pdf_and_docx_encrypted_artifacts(tmp_path: Path) -> None:
    client, store, queue, worker = _build_test_stack(tmp_path)
    secret_marker = "CONFIDENTIAL_MARKER_998877"
    pdf_bytes = _make_sample_pdf(secret_marker=secret_marker)

    submit_resp = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "tenant_01"},
        data={"task_type": "translate", "target_language": "Spanish"},
        files={"file": ("contract.pdf", pdf_bytes, "application/pdf")},
    )
    assert submit_resp.status_code == 202
    job_id = submit_resp.json()["job_id"]
    assert queue.queue_depth() == 1

    # Verify ciphertext on disk does NOT expose plaintext PDF marker
    enc_files = list((tmp_path / "artifacts").rglob("*.enc"))
    assert len(enc_files) == 1
    raw_ciphertext = enc_files[0].read_bytes()
    assert b"%PDF-" not in raw_ciphertext
    assert secret_marker.encode("utf-8") not in raw_ciphertext

    # Artifact download before worker completion returns 409 Conflict
    early_dl = client.get(
        f"/api/v1/jobs/{job_id}/artifacts/pdf",
        headers={"X-Session-ID": "tenant_01"},
    )
    assert early_dl.status_code == 409

    # Run async worker
    processed = worker.process_next()
    assert processed is not None
    assert processed.state == JobState.COMPLETED
    assert processed.current_stage == JobStage.VERIFIED
    assert queue.queue_depth() == 0

    status_resp = client.get(f"/api/v1/jobs/{job_id}", headers={"X-Session-ID": "tenant_01"})
    assert status_resp.status_code == 200
    status_body = status_resp.json()
    assert status_body["state"] == "completed"
    assert status_body["traffic_light"] in {"green", "yellow", "red"}
    assert set(status_body["download_urls"].keys()) == {"pdf", "docx", "report"}

    # Download decrypted PDF, DOCX, and JSON report
    pdf_dl = client.get(status_body["download_urls"]["pdf"], headers={"X-Session-ID": "tenant_01"})
    assert pdf_dl.status_code == 200
    assert pdf_dl.content.startswith(b"%PDF-")

    docx_dl = client.get(
        status_body["download_urls"]["docx"], headers={"X-Session-ID": "tenant_01"}
    )
    assert docx_dl.status_code == 200
    assert docx_dl.content.startswith(b"PK\x03\x04")

    rep_dl = client.get(
        status_body["download_urls"]["report"], headers={"X-Session-ID": "tenant_01"}
    )
    assert rep_dl.status_code == 200
    parsed_rep = json.loads(rep_dl.content.decode("utf-8"))
    assert parsed_rep["job_id"] == job_id
    assert "quality_report" in parsed_rep

    # Also verify DOCX source pipeline via ZeroQuotaSimulationClient
    worker_sim = AsyncDocumentWorker(
        store=store,
        queue=queue,
        llm_client=ZeroQuotaSimulationClient(),
    )
    docx_submit = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "tenant_01"},
        data={"task_type": "rewrite", "target_language": "Spanish"},
        files={"file": ("memo.docx", _make_sample_docx(), "application/octet-stream")},
    )
    assert docx_submit.status_code == 202
    rec_docx = worker_sim.process_next()
    assert rec_docx is not None
    assert rec_docx.state == JobState.COMPLETED


def test_idempotency_replay_and_concurrency_limiting(tmp_path: Path) -> None:
    client, _, queue, _ = _build_test_stack(tmp_path, max_concurrent=1)
    pdf_bytes = _make_sample_pdf()

    r1 = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_idem", "Idempotency-Key": "req-key-0001"},
        files={"file": ("doc1.pdf", pdf_bytes, "application/pdf")},
    )
    assert r1.status_code == 202
    assert r1.json()["idempotent_replay"] is False

    # Duplicate submission with same Idempotency-Key replays existing job without 429 or re-enqueue
    r2 = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_idem", "Idempotency-Key": "req-key-0001"},
        files={"file": ("doc1.pdf", pdf_bytes, "application/pdf")},
    )
    assert r2.status_code == 202
    assert r2.json()["job_id"] == r1.json()["job_id"]
    assert r2.json()["idempotent_replay"] is True
    assert queue.queue_depth() == 1

    # Distinct Idempotency-Key while 1 active job is already queued hits 429 concurrency limit
    r3 = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "session_idem", "Idempotency-Key": "req-key-0002"},
        files={"file": ("doc2.pdf", pdf_bytes, "application/pdf")},
    )
    assert r3.status_code == 429


def test_strict_session_isolation_and_delete_shredding(tmp_path: Path) -> None:
    client, _, _, _ = _build_test_stack(tmp_path)
    resp = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "alice_session"},
        files={"file": ("private.pdf", _make_sample_pdf(), "application/pdf")},
    )
    job_id = resp.json()["job_id"]

    # Bob cannot list, inspect, download, or delete Alice's job
    assert client.get("/api/v1/jobs", headers={"X-Session-ID": "bob_session"}).json() == []
    assert (
        client.get(f"/api/v1/jobs/{job_id}", headers={"X-Session-ID": "bob_session"}).status_code
        == 404
    )
    assert (
        client.delete(f"/api/v1/jobs/{job_id}", headers={"X-Session-ID": "bob_session"}).status_code
        == 404
    )

    # Alice deletes her job -> encrypted directory is shredded immediately
    del_resp = client.delete(f"/api/v1/jobs/{job_id}", headers={"X-Session-ID": "alice_session"})
    assert del_resp.status_code == 204
    assert list((tmp_path / "artifacts").rglob("*.enc")) == []
    assert (
        client.get(f"/api/v1/jobs/{job_id}", headers={"X-Session-ID": "alice_session"}).status_code
        == 404
    )


def test_worker_checkpoint_recovery_skips_duplicate_llm_call(tmp_path: Path) -> None:
    class _CountingLLMClient:
        def __init__(self) -> None:
            self.calls = 0
            self._delegate = MockGeminiClient()

        def transform_document(
            self,
            layout: DocumentLayout,
            request: TransformationRequest,
        ) -> DocumentTransformationOutput:
            self.calls += 1
            return self._delegate.transform_document(layout=layout, request=request)

    client, store, queue, _ = _build_test_stack(tmp_path)
    counting_llm = _CountingLLMClient()
    worker = AsyncDocumentWorker(store=store, queue=queue, llm_client=counting_llm)

    pdf_bytes = _make_sample_pdf()
    resp = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "ckpt_user"},
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")},
    )
    job_id = resp.json()["job_id"]

    # First run populates checkpoints and completes
    first = worker.process_next()
    assert first is not None and first.state == JobState.COMPLETED
    assert counting_llm.calls == 1

    # Simulate interrupted worker retry before completion: reset state to QUEUED while keeping ckpts
    rec = store.get_job("ckpt_user", job_id)
    assert rec is not None
    rec.state = JobState.QUEUED
    rec.current_stage = JobStage.TRANSFORMED
    store.update_job(rec)
    queue.enqueue(job_id=job_id, session_id="ckpt_user")

    resumed = worker.process_next()
    assert resumed is not None
    assert resumed.state == JobState.COMPLETED
    # Crucial: LLM was NOT called a second time because encrypted checkpoint existed!
    assert counting_llm.calls == 1


def test_ttl_expiration_and_rate_limiter_and_probe_server(tmp_path: Path) -> None:
    client, store, queue, worker = _build_test_stack(tmp_path, ttl_seconds=60)
    resp = client.post(
        "/api/v1/jobs",
        headers={"X-Session-ID": "ttl_user"},
        files={"file": ("short_lived.pdf", _make_sample_pdf(), "application/pdf")},
    )
    assert resp.status_code == 202
    assert len(list((tmp_path / "artifacts").rglob("*.enc"))) == 1

    future = datetime.now(UTC) + timedelta(seconds=120)
    purged = store.purge_expired(now=future)
    assert purged == 1
    assert list((tmp_path / "artifacts").rglob("*.enc")) == []

    # SlidingWindowRateLimiter test
    limiter = SlidingWindowRateLimiter(max_requests_per_minute=2, window_seconds=0.15)
    assert limiter.try_acquire()[0] is True
    assert limiter.try_acquire()[0] is True
    allowed, retry_after = limiter.try_acquire()
    assert allowed is False
    assert retry_after > 0.0
    assert limiter.acquire_blocking(max_wait_seconds=0.5) is True

    # WorkerProbeServer live HTTP test on ephemeral port
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))
        free_port = int(s.getsockname()[1])

    probe = WorkerProbeServer(worker=worker, host="127.0.0.1", port=free_port)
    probe.start()
    try:
        with urllib.request.urlopen(f"http://127.0.0.1:{free_port}/healthz", timeout=3) as r_h:
            assert r_h.status == 200
        with urllib.request.urlopen(f"http://127.0.0.1:{free_port}/readyz", timeout=3) as r_r:
            assert r_r.status == 200

        worker.request_stop()
        try:
            urllib.request.urlopen(f"http://127.0.0.1:{free_port}/readyz", timeout=3)
            raise AssertionError("Expected 503 when worker is draining")
        except urllib.error.HTTPError as http_err:
            assert http_err.code == 503
    finally:
        probe.stop()


def test_kubernetes_manifests_security_and_architecture_audit() -> None:
    base_dir = REPO_ROOT / "k8s" / "base"
    assert base_dir.is_dir()

    docs: list[dict[str, Any]] = []
    for yml_path in sorted(base_dir.glob("*.yaml")):
        raw_docs = list(yaml.safe_load_all(yml_path.read_text(encoding="utf-8")))
        docs.extend(d for d in raw_docs if isinstance(d, dict))

    deployments = {d["metadata"]["name"]: d for d in docs if d.get("kind") == "Deployment"}
    assert {"doclayout-ui", "doclayout-api", "doclayout-worker", "doclayout-redis"}.issubset(
        deployments.keys()
    )

    for dep_name, dep in deployments.items():
        pod_spec = dep["spec"]["template"]["spec"]
        assert pod_spec.get("automountServiceAccountToken") is False, dep_name
        pod_sec = pod_spec.get("securityContext", {})
        assert pod_sec.get("runAsNonRoot") is True, dep_name
        assert int(pod_sec.get("runAsUser", 0)) >= 10000, dep_name

        for container in pod_spec["containers"]:
            c_sec = container.get("securityContext", {})
            assert c_sec.get("runAsNonRoot") is True, f"{dep_name}/{container['name']}"
            assert c_sec.get("readOnlyRootFilesystem") is True, f"{dep_name}/{container['name']}"
            assert c_sec.get("allowPrivilegeEscalation") is False, f"{dep_name}/{container['name']}"
            assert c_sec.get("capabilities", {}).get("drop") == ["ALL"]

            res = container.get("resources", {})
            assert "requests" in res and "limits" in res
            assert "startupProbe" in container
            assert "readinessProbe" in container
            assert "livenessProbe" in container
            assert not str(container.get("image", "")).endswith(":latest")

    # Verify HPA includes queue depth + CPU + Memory
    hpas = {d["metadata"]["name"]: d for d in docs if d.get("kind") == "HorizontalPodAutoscaler"}
    assert "doclayout-worker-hpa" in hpas
    worker_metrics = hpas["doclayout-worker-hpa"]["spec"]["metrics"]
    metric_names = {
        m.get("external", {}).get("metric", {}).get("name") or m.get("resource", {}).get("name")
        for m in worker_metrics
    }
    assert {"doclayout_queue_depth", "cpu", "memory"}.issubset(metric_names)

    # Verify PDBs exist for UI, API, and Worker
    pdbs = {d["metadata"]["name"] for d in docs if d.get("kind") == "PodDisruptionBudget"}
    assert {"doclayout-ui-pdb", "doclayout-api-pdb", "doclayout-worker-pdb"}.issubset(pdbs)

    # Verify default deny-all NetworkPolicy
    netpols = {d["metadata"]["name"]: d for d in docs if d.get("kind") == "NetworkPolicy"}
    assert "default-deny-all" in netpols
    deny_spec = netpols["default-deny-all"]["spec"]
    assert deny_spec["podSelector"] == {}
    assert set(deny_spec["policyTypes"]) == {"Ingress", "Egress"}

    # Verify validation & production overlays
    for env_name in ("validation", "production"):
        overlay_file = REPO_ROOT / "k8s" / "overlays" / env_name / "kustomization.yaml"
        assert overlay_file.is_file()
        overlay_cfg = yaml.safe_load(overlay_file.read_text(encoding="utf-8"))
        assert "../../base" in overlay_cfg["resources"]
