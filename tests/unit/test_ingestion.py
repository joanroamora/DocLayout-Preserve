"""Unit tests for PDF and DOCX layout ingestion and structural extraction."""

from __future__ import annotations

import io

import docx
import pymupdf
import pytest

from doclayout.core.models import BlockType
from doclayout.ingestion import (
    CorruptedDocumentError,
    DocumentLimitExceededError,
    DocxLayoutExtractor,
    EmptyDocumentError,
    PasswordProtectedDocumentError,
    PDFLayoutExtractor,
    extract_document_layout,
)


def _build_sample_pdf_bytes() -> bytes:
    """Create an in-memory PDF with header, heading, paragraph, list, and bordered table."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)

    # Header at very top
    page.insert_text((50, 30), "CONFIDENTIAL HEADER", fontsize=9)
    # Heading (fontsize >= 14)
    page.insert_text((50, 90), "1. Executive Architecture Summary", fontsize=16)
    # Paragraph
    page.insert_text(
        (50, 140),
        "This document describes the layout preservation pipeline and coordinates.",
        fontsize=11,
    )
    # List item
    page.insert_text((50, 190), "- First bullet item in the specification", fontsize=11)

    # Draw a 2x2 table grid so pdfplumber detects a table
    x_cols = [50, 200, 350]
    y_rows = [260, 290, 320]
    for y in y_rows:
        page.draw_line((x_cols[0], y), (x_cols[-1], y))
    for x in x_cols:
        page.draw_line((x, y_rows[0]), (x, y_rows[-1]))

    page.insert_text((60, 280), "Metric", fontsize=10)
    page.insert_text((210, 280), "Target", fontsize=10)
    page.insert_text((60, 310), "IoU", fontsize=10)
    page.insert_text((210, 310), "0.95", fontsize=10)

    # Footer at bottom
    page.insert_text((50, 810), "Page 1 of 1", fontsize=9)

    data = bytes(doc.tobytes())
    doc.close()
    return data


def _build_sample_docx_bytes() -> bytes:
    """Create an in-memory Word .docx with heading, paragraph, list, and table."""
    doc = docx.Document()
    doc.add_heading("Technical Architecture Document", level=1)
    doc.add_paragraph("Standard body paragraph preserving typographic metadata.")
    doc.add_paragraph("- Structured bullet item for testing.")

    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Component"
    tbl.cell(0, 1).text = "Role"
    tbl.cell(1, 0).text = "Extractor"
    tbl.cell(1, 1).text = "Bounding Boxes"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_layout_extractor_blocks_and_tables() -> None:
    """Verify PDF extraction captures headers, headings, paragraphs, lists, tables, and footers."""
    pdf_bytes = _build_sample_pdf_bytes()
    layout = PDFLayoutExtractor().extract_from_bytes(pdf_bytes, filename="spec.pdf")

    assert layout.source_format == "pdf"
    assert len(layout.pages) == 1
    blocks = layout.all_blocks
    block_types = {b.block_type for b in blocks}

    assert BlockType.HEADER in block_types
    assert BlockType.HEADING in block_types
    assert BlockType.PARAGRAPH in block_types
    assert BlockType.LIST in block_types
    assert BlockType.TABLE in block_types
    assert BlockType.FOOTER in block_types

    table_blocks = [b for b in blocks if b.block_type == BlockType.TABLE]
    assert len(table_blocks) == 1
    assert len(table_blocks[0].cells) == 4
    assert table_blocks[0].metadata["col_count"] == 2
    assert table_blocks[0].metadata["relative_col_widths"] == [0.5, 0.5]


def test_pdf_password_protected_raises_error() -> None:
    """Verify encrypted PDF raises PasswordProtectedDocumentError."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 100), "Secret text")
    encrypted_bytes = doc.tobytes(
        encryption=int(getattr(pymupdf, "PDF_ENCRYPT_AES_256", 4)),
        user_pw="secret123",
        owner_pw="owner123",
    )
    doc.close()

    with pytest.raises(PasswordProtectedDocumentError):
        PDFLayoutExtractor().extract_from_bytes(encrypted_bytes, filename="locked.pdf")


def test_pdf_empty_and_corrupted_errors() -> None:
    """Verify empty, blank-page, or corrupted PDF bytes raise appropriate exceptions."""
    extractor = PDFLayoutExtractor()

    with pytest.raises(EmptyDocumentError):
        extractor.extract_from_bytes(b"", filename="empty.pdf")

    with pytest.raises(CorruptedDocumentError):
        extractor.extract_from_bytes(b"not-a-valid-pdf-stream", filename="bad.pdf")

    blank_doc = pymupdf.open()
    blank_doc.new_page()
    blank_bytes = blank_doc.tobytes()
    blank_doc.close()

    with pytest.raises(EmptyDocumentError):
        extractor.extract_from_bytes(blank_bytes, filename="blank.pdf")


def test_pdf_page_limit_exceeded() -> None:
    """Verify PDF exceeding max_pages raises DocumentLimitExceededError."""
    doc = pymupdf.open()
    for i in range(3):
        p = doc.new_page()
        p.insert_text((50, 100), f"Page {i}")
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(DocumentLimitExceededError):
        PDFLayoutExtractor(max_pages=2).extract_from_bytes(pdf_bytes, filename="long.pdf")


def test_docx_layout_extractor_and_dispatcher() -> None:
    """Verify DOCX extraction captures headings, paragraphs, lists, and tables."""
    docx_bytes = _build_sample_docx_bytes()
    layout = extract_document_layout(docx_bytes, filename="sample.docx")

    assert layout.source_format == "docx"
    assert len(layout.pages) >= 1
    blocks = layout.all_blocks
    block_types = {b.block_type for b in blocks}

    assert BlockType.HEADING in block_types
    assert BlockType.PARAGRAPH in block_types
    assert BlockType.LIST in block_types
    assert BlockType.TABLE in block_types

    tbl_block = next(b for b in blocks if b.block_type == BlockType.TABLE)
    assert len(tbl_block.cells) == 4
    assert tbl_block.metadata["relative_col_widths"] == [0.5, 0.5]


def test_docx_empty_and_unsupported_extension() -> None:
    """Verify empty DOCX and unsupported file extensions raise expected exceptions."""
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)

    with pytest.raises(EmptyDocumentError):
        DocxLayoutExtractor().extract_from_bytes(buf.getvalue(), filename="empty.docx")

    with pytest.raises(CorruptedDocumentError):
        extract_document_layout(b"hello", filename="unsupported.txt")
